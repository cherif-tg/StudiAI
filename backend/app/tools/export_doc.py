import io
import re
from xml.sax.saxutils import escape

from app.tools.base import Tool, ToolContext, ToolResult


def parse_markdown(md: str) -> list[tuple[str, str]]:
    blocks = []
    for line in md.splitlines():
        line = line.rstrip()
        if not line.strip():
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", line)
        if m:
            blocks.append((f"h{len(m.group(1))}", m.group(2)))
        elif re.match(r"^\s*[-*]\s+", line):
            blocks.append(("li", re.sub(r"^\s*[-*]\s+", "", line)))
        else:
            blocks.append(("p", line))
    return blocks


def strip_inline(text: str) -> str:
    return re.sub(r"[*_`]{1,3}", "", text)


def to_pdf(title: str, md: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    story = [Paragraph(escape(title), styles["Title"]), Spacer(1, 12)]
    for kind, text in parse_markdown(md):
        text = escape(strip_inline(text))
        if kind in {"h1", "h2", "h3"}:
            story.append(Paragraph(text, styles[{"h1": "Heading1", "h2": "Heading2", "h3": "Heading3"}[kind]]))
        elif kind == "li":
            story.append(Paragraph(f"• {text}", styles["BodyText"]))
        else:
            story.append(Paragraph(text, styles["BodyText"]))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm).build(story)
    return buf.getvalue()


def to_docx(title: str, md: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(title, 0)
    for kind, text in parse_markdown(md):
        text = strip_inline(text)
        if kind in {"h1", "h2", "h3"}:
            doc.add_heading(text, int(kind[1]))
        elif kind == "li":
            doc.add_paragraph(text, style="List Bullet")
        else:
            doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class ExportDocTool(Tool):
    name = "export_document"
    description = "Exporte un contenu (fiche de révision, résumé…) en fichier PDF ou Word (DOCX) téléchargeable."
    parameters = {
        "type": "object",
        "properties": {
            "title": {"type": "string"},
            "content": {"type": "string", "description": "Contenu en Markdown (titres #, listes -)"},
            "format": {"type": "string", "enum": ["pdf", "docx"]},
        },
        "required": ["title", "content", "format"],
    }

    async def run(self, ctx: ToolContext, title: str, content: str, format: str = "pdf") -> ToolResult:
        if format == "docx":
            data, ext, mime = to_docx(title, content), "docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        else:
            data, ext, mime = to_pdf(title, content), "pdf", "application/pdf"
        art = ctx.save_artifact(ext, title, data, ext, mime)
        return ToolResult(f"Document « {title} » exporté en {ext.upper()}.", artifacts=[art])
