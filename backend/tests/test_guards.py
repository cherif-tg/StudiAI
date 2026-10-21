"""Garde-fous d'exposition publique : taille, stockage, débit."""
import io

from docx import Document

from app.services.container import build_services, get_services


def _docx(remplissage: int = 0) -> bytes:
    document = Document()
    document.add_paragraph("La photosynthèse convertit l'énergie lumineuse.")
    if remplissage:
        document.add_paragraph("x" * remplissage)
    tampon = io.BytesIO()
    document.save(tampon)
    return tampon.getvalue()


def _appliquer(client, settings, **reglages):
    """Remplace les services par une variante aux garde-fous resserrés.

    Le dossier de données reste le même : la session ouverte par `connecte`
    demeure valide.
    """
    base = settings.model_copy(update=reglages)
    client.app.dependency_overrides[get_services] = lambda: build_services(base)
    return base


def test_upload_trop_volumineux_refuse_avec_413(connecte, settings):
    _appliquer(connecte, settings, max_upload_mb=1)

    contenu = b"x" * (2 * 1024 * 1024)  # 2 Mo pour un plafond de 1 Mo
    r = connecte.post("/sources", files={"file": ("gros.docx", contenu, "application/octet-stream")})

    assert r.status_code == 413
    assert "maximum 1 Mo" in r.json()["detail"]


def test_stockage_plein_refuse_le_depot(connecte, settings):
    base = _appliquer(connecte, settings, max_upload_mb=20, max_storage_mb=1)

    remplissage = base.uploads_dir / "deja-la.bin"
    remplissage.parent.mkdir(parents=True, exist_ok=True)
    remplissage.write_bytes(b"y" * (2 * 1024 * 1024))
    try:
        r = connecte.post(
            "/sources", files={"file": ("cours.docx", _docx(), "application/octet-stream")}
        )
        assert r.status_code == 507
        assert "stockage plein" in r.json()["detail"].lower()
    finally:
        remplissage.unlink(missing_ok=True)


def test_debit_sur_chat_limite_par_client(connecte, settings):
    _appliquer(connecte, settings, chat_rate_limit=3, chat_rate_window_s=600)
    entetes = {"cf-connecting-ip": "203.0.113.7"}

    for _ in range(3):
        r = connecte.post(
            "/chat", json={"session_id": "s-debit", "message": "bonjour"}, headers=entetes
        )
        assert r.status_code == 503  # pas de clé LLM : la requête franchit le limiteur

    bloquee = connecte.post(
        "/chat", json={"session_id": "s-debit", "message": "encore"}, headers=entetes
    )

    assert bloquee.status_code == 429
    assert "Retry-After" in bloquee.headers


def test_debit_isole_par_adresse(connecte, settings):
    _appliquer(connecte, settings, chat_rate_limit=1, chat_rate_window_s=600)

    premier = connecte.post(
        "/chat",
        json={"session_id": "s-a", "message": "a"},
        headers={"cf-connecting-ip": "198.51.100.1"},
    )
    autre = connecte.post(
        "/chat",
        json={"session_id": "s-b", "message": "b"},
        headers={"cf-connecting-ip": "198.51.100.2"},
    )

    # Le second client ne doit pas être pénalisé par le premier.
    assert premier.status_code == 503
    assert autre.status_code == 503


def test_suppression_retire_le_fichier_depose(connecte, settings, monkeypatch):
    from app.services.rag import RagService

    _appliquer(connecte, settings)
    monkeypatch.setattr(RagService, "ingest", lambda self, source_id, path, filename: 2)
    monkeypatch.setattr(RagService, "delete", lambda self, source_id: None)

    depose = connecte.post(
        "/sources", files={"file": ("a-retirer.docx", _docx(), "application/octet-stream")}
    )
    assert depose.status_code == 200

    identifiant = depose.json()["id"]
    connecte.delete(f"/sources/{identifiant}")

    assert [s for s in connecte.get("/sources").json() if s["id"] == identifiant] == []
