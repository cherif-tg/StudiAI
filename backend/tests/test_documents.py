"""Extraction de documents : le contenu des tableaux doit être indexé."""
from pathlib import Path

from docx import Document

from app.services.documents import chunk_text, extract_text


def _docx_avec_tableau(path: Path) -> Path:
    doc = Document()
    doc.add_heading("Cours de biologie", level=1)
    doc.add_paragraph("La photosynthèse convertit l'énergie lumineuse en énergie chimique.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Phase"
    table.cell(0, 1).text = "Lieu"
    table.cell(1, 0).text = "Phase claire"
    table.cell(1, 1).text = "Thylakoïdes"
    doc.save(path)
    return path


def test_extract_text_lit_les_tableaux(tmp_path: Path):
    texte = extract_text(_docx_avec_tableau(tmp_path / "cours.docx"))

    assert "photosynthèse" in texte
    # Sans extraction des tableaux, ces cellules seraient absentes de l'index
    # et donc introuvables par le RAG.
    assert "Phase claire" in texte
    assert "Thylakoïdes" in texte


def test_extract_text_refuse_un_format_inconnu(tmp_path: Path):
    fichier = tmp_path / "notes.odt"
    fichier.write_text("contenu", encoding="utf-8")

    try:
        extract_text(fichier)
    except ValueError as exc:
        assert "Format non supporté" in str(exc)
    else:  # pragma: no cover
        raise AssertionError("un format inconnu doit lever ValueError")


def test_chunk_text_respecte_la_taille_et_conserve_le_texte():
    paragraphe = " ".join(f"mot{i}" for i in range(400))
    chunks = chunk_text(paragraphe, size=1000, overlap=150)

    assert len(chunks) >= 2
    assert all(len(c) <= 1200 for c in chunks)
    assert "mot0" in chunks[0]
    assert "mot399" in chunks[-1]


def test_chunk_text_sur_texte_vide():
    assert chunk_text("   \n\n  ") == []
